# app/api/reports.py - COMPLETE WITH JSON & EXCEL EXPORTS
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, JSONResponse
from sqlalchemy.orm import Session
from ..database import get_db
from ..models.session import ValidationSession
from ..models.swab_result import SwabResult
from ..models.rinse_result import RinseResult
from ..models.session_equipment import SessionEquipment
from ..models.product import Product
from ..models.equipment import Equipment
from ..services.report import ReportService
from ..api.dependencies import get_current_user
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter()

@router.get("/{session_id}/pdf")
def generate_report(
    session_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    session = db.query(ValidationSession).filter(ValidationSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    
    swab_results = db.query(SwabResult).filter(SwabResult.session_id == session_id).all()
    rinse_results = db.query(RinseResult).filter(RinseResult.session_id == session_id).all()
    
    session_equipment = db.query(SessionEquipment).filter(SessionEquipment.session_id == session_id).all()
    equipment_list = [se.equipment for se in session_equipment if se.equipment]
    
    maco_data = {
        "10ppm": session.maco_10ppm,
        "tdd": session.maco_tdd,
        "ade_pde": session.maco_ade_pde,
        "lowest": session.lowest_maco
    }
    
    pdf_content = ReportService.generate_validation_report(
        session, swab_results, rinse_results, maco_data, equipment_list
    )
    
    return Response(
        content=pdf_content,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=validation_report_{session.session_code}.pdf"}
    )


@router.get("/{session_id}/json")
def export_json(
    session_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Export validation data as JSON"""
    session = db.query(ValidationSession).filter(ValidationSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    
    swab_results = db.query(SwabResult).filter(SwabResult.session_id == session_id).all()
    rinse_results = db.query(RinseResult).filter(RinseResult.session_id == session_id).all()
    
    data = {
        "session": {
            "id": session.id,
            "session_code": session.session_code,
            "status": session.status,
            "created_at": session.created_at.isoformat() if session.created_at else None,
            "extra_area_percentage": session.extra_area_percentage,
            "total_surface_area": session.total_surface_area
        },
        "previous_product": None,
        "next_product": None,
        "maco_calculations": {
            "method_10ppm": session.maco_10ppm,
            "method_tdd": session.maco_tdd,
            "method_ade_pde": session.maco_ade_pde,
            "lowest_maco": session.lowest_maco
        },
        "limits": {
            "swab_limit_mg": session.swab_limit_mg,
            "swab_limit_ppm": session.swab_limit_ppm,
            "rinse_limit_mg": session.rinse_limit_mg,
            "rinse_limit_ppm": session.rinse_limit_ppm
        },
        "swab_results": [
            {
                "location_name": r.location_name,
                "absorbance_sample": r.absorbance_sample,
                "absorbance_std": r.absorbance_std,
                "result_mg_ml": r.result_mg_ml,
                "result_ppm": r.result_ppm,
                "reported": r.reported
            }
            for r in swab_results
        ],
        "rinse_results": [
            {
                "equipment_name": r.equipment_name,
                "actual_rinse_volume": r.actual_rinse_volume,
                "absorbance_sample": r.absorbance_sample,
                "absorbance_std": r.absorbance_std,
                "result_mg_ml": r.result_mg_ml,
                "result_ppm": r.result_ppm,
                "reported": r.reported
            }
            for r in rinse_results
        ]
    }
    
    if session.previous_product:
        data["previous_product"] = {
            "id": session.previous_product.id,
            "name": session.previous_product.name,
            "min_batch_size": session.previous_product.min_batch_size,
            "max_batch_size": session.previous_product.max_batch_size,
            "ade_pde": session.previous_product.ade_pde
        }
    
    if session.next_product:
        data["next_product"] = {
            "id": session.next_product.id,
            "name": session.next_product.name,
            "min_batch_size": session.next_product.min_batch_size,
            "max_batch_size": session.next_product.max_batch_size,
            "solubility": session.next_product.solubility
        }
    
    return JSONResponse(content=data)


@router.get("/{session_id}/excel")
def export_excel(
    session_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Export validation data as Excel"""
    session = db.query(ValidationSession).filter(ValidationSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    
    swab_results = db.query(SwabResult).filter(SwabResult.session_id == session_id).all()
    rinse_results = db.query(RinseResult).filter(RinseResult.session_id == session_id).all()
    
    workbook = openpyxl.Workbook()
    
    # Session Info Sheet
    info_sheet = workbook.active
    info_sheet.title = "Session Info"
    info_sheet.cell(1, 1, "Session Code")
    info_sheet.cell(1, 2, session.session_code)
    info_sheet.cell(2, 1, "Status")
    info_sheet.cell(2, 2, session.status)
    info_sheet.cell(3, 1, "Created At")
    info_sheet.cell(3, 2, session.created_at.isoformat() if session.created_at else "")
    info_sheet.cell(4, 1, "Lowest MACO (mg)")
    info_sheet.cell(4, 2, session.lowest_maco)
    info_sheet.cell(5, 1, "Swab Limit (ppm)")
    info_sheet.cell(5, 2, session.swab_limit_ppm)
    info_sheet.cell(6, 1, "Rinse Limit (ppm)")
    info_sheet.cell(6, 2, session.rinse_limit_ppm)
    
    # Swab Results Sheet
    swab_sheet = workbook.create_sheet("Swab Results")
    headers = ["Location", "Abs Sample", "Abs Std", "Result mg/ml", "Result ppm", "Reported"]
    for col, header in enumerate(headers, 1):
        cell = swab_sheet.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="2d6a4f", end_color="2d6a4f", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
    
    for idx, result in enumerate(swab_results, 2):
        swab_sheet.cell(idx, 1, result.location_name)
        swab_sheet.cell(idx, 2, result.absorbance_sample)
        swab_sheet.cell(idx, 3, result.absorbance_std)
        swab_sheet.cell(idx, 4, result.result_mg_ml)
        swab_sheet.cell(idx, 5, result.result_ppm)
        swab_sheet.cell(idx, 6, result.reported)
    
    # Rinse Results Sheet
    rinse_sheet = workbook.create_sheet("Rinse Results")
    rinse_headers = ["Equipment", "Rinse Volume (L)", "Abs Sample", "Abs Std", "Result mg/ml", "Result ppm", "Reported"]
    for col, header in enumerate(rinse_headers, 1):
        cell = rinse_sheet.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="2d6a4f", end_color="2d6a4f", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
    
    for idx, result in enumerate(rinse_results, 2):
        rinse_sheet.cell(idx, 1, result.equipment_name)
        rinse_sheet.cell(idx, 2, result.actual_rinse_volume)
        rinse_sheet.cell(idx, 3, result.absorbance_sample)
        rinse_sheet.cell(idx, 4, result.absorbance_std)
        rinse_sheet.cell(idx, 5, result.result_mg_ml)
        rinse_sheet.cell(idx, 6, result.result_ppm)
        rinse_sheet.cell(idx, 7, result.reported)
    
    # Auto-adjust column widths
    for sheet in [info_sheet, swab_sheet, rinse_sheet]:
        for column in sheet.columns:
            max_length = 0
            column_letter = openpyxl.utils.get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=validation_data_{session.session_code}.xlsx"}
    )