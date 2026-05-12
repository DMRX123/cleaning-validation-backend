import openpyxl
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from ..models.product import Product
from ..models.equipment import Equipment
from ..models.session import ValidationSession
from ..models.swab_result import SwabResult
from ..models.rinse_result import RinseResult

def export_products_to_excel(db: Session) -> BytesIO:
    """Export all products to Excel in original format"""
    
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Input Details"
    
    # Headers (matching original Excel)
    headers = [
        "Sr. No.", "Product Name", "Min Batch Size (kg)", "Max Batch Size (kg)",
        "ADE/PDE (µg/day)", "Min Dose (mg)", "Max Dose (mg)", "Swab Recovery %",
        "LOD in ppm", "LOQ in ppm", "Swab Dilution in ml", "Swab Surface in M. Sq.",
        "Min Batch/Max Dose Ratio", "Solubility", "Hardest To Clean", "Plant / Block Name"
    ]
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2d6a4f", end_color="2d6a4f", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    products = db.query(Product).all()
    
    for idx, product in enumerate(products, 1):
        row = idx + 1
        sheet.cell(row, 1, idx)
        sheet.cell(row, 2, product.name)
        sheet.cell(row, 3, product.min_batch_size)
        sheet.cell(row, 4, product.max_batch_size)
        sheet.cell(row, 5, product.ade_pde)
        sheet.cell(row, 6, product.min_dose)
        sheet.cell(row, 7, product.max_dose)
        sheet.cell(row, 8, product.swab_recovery)
        sheet.cell(row, 9, product.lod)
        sheet.cell(row, 10, product.loq)
        sheet.cell(row, 11, product.swab_dilution)
        sheet.cell(row, 12, product.swab_surface_area)
        sheet.cell(row, 13, product.get_min_batch_max_dose_ratio())
        sheet.cell(row, 14, product.solubility)
        sheet.cell(row, 15, product.hardest_to_clean)
        sheet.cell(row, 16, product.plant)
    
    # Auto-adjust column widths
    for col in range(1, 17):
        max_length = 0
        for row in range(1, len(products) + 2):
            cell_value = sheet.cell(row, col).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max_length + 2, 30)
    
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def export_results_to_excel(db: Session, session_id: int) -> BytesIO:
    """Export validation results for a session"""
    
    workbook = openpyxl.Workbook()
    
    # Session Info Sheet
    session_sheet = workbook.active
    session_sheet.title = "Session Info"
    
    session = db.query(ValidationSession).filter(ValidationSession.id == session_id).first()
    if session:
        session_sheet.cell(1, 1, "Session Code")
        session_sheet.cell(1, 2, session.session_code)
        session_sheet.cell(2, 1, "Previous Product")
        session_sheet.cell(2, 2, session.previous_product.name if session.previous_product else "N/A")
        session_sheet.cell(3, 1, "Next Product")
        session_sheet.cell(3, 2, session.next_product.name if session.next_product else "N/A")
        session_sheet.cell(4, 1, "Lowest MACO (mg)")
        session_sheet.cell(4, 2, session.lowest_maco)
        session_sheet.cell(5, 1, "Swab Limit (ppm)")
        session_sheet.cell(5, 2, session.swab_limit_ppm)
        session_sheet.cell(6, 1, "Rinse Limit (ppm)")
        session_sheet.cell(6, 2, session.rinse_limit_ppm)
    
    # Swab Results Sheet
    swab_sheet = workbook.create_sheet("Swab Results")
    swab_headers = ["Location", "Abs Sample", "Abs Std", "Result mg/ml", "Result ppm", "Reported"]
    for col, header in enumerate(swab_headers, 1):
        swab_sheet.cell(1, col, header)
    
    swab_results = db.query(SwabResult).filter(SwabResult.session_id == session_id).all()
    for idx, result in enumerate(swab_results, 2):
        swab_sheet.cell(idx, 1, result.location_name)
        swab_sheet.cell(idx, 2, result.absorbance_sample)
        swab_sheet.cell(idx, 3, result.absorbance_std)
        swab_sheet.cell(idx, 4, result.result_mg_ml)
        swab_sheet.cell(idx, 5, result.result_ppm)
        swab_sheet.cell(idx, 6, result.reported)
    
    # Rinse Results Sheet
    rinse_sheet = workbook.create_sheet("Rinse Results")
    rinse_headers = ["Equipment", "Rinse Volume (L)", "Abs Sample", "Abs Std", "Result mg/ml", "Result ppm", "Reported"]
    for col, header in enumerate(rinse_headers, 1):
        rinse_sheet.cell(1, col, header)
    
    rinse_results = db.query(RinseResult).filter(RinseResult.session_id == session_id).all()
    for idx, result in enumerate(rinse_results, 2):
        rinse_sheet.cell(idx, 1, result.equipment_name)
        rinse_sheet.cell(idx, 2, result.actual_rinse_volume)
        rinse_sheet.cell(idx, 3, result.absorbance_sample)
        rinse_sheet.cell(idx, 4, result.absorbance_std)
        rinse_sheet.cell(idx, 5, result.result_mg_ml)
        rinse_sheet.cell(idx, 6, result.result_ppm)
        rinse_sheet.cell(idx, 7, result.reported)
    
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer