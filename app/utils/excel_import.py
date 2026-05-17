import openpyxl
import io
from fastapi import UploadFile
from sqlalchemy.orm import Session
from ..models.product import Product
from ..models.equipment import Equipment


def import_products_from_excel(file: UploadFile, db: Session) -> dict:
    """Import products from Excel file matching the original structure (SYNC VERSION)"""
    
    contents = file.file.read()
    workbook = openpyxl.load_workbook(io.BytesIO(contents))
    
    # Try to find the correct sheet
    sheet = None
    for sheet_name in ["Input Details", "Sheet1", "Products"]:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            break
    
    if not sheet:
        return {"error": "Could not find data sheet. Expected 'Input Details'"}
    
    imported_count = 0
    skipped_count = 0
    errors = []
    
    # Start from row 4 (data starts in original Excel)
    for row in range(4, sheet.max_row + 1):
        product_name = sheet.cell(row, 3).value  # Column C
        if not product_name:
            continue
        
        # Check if product already exists
        existing = db.query(Product).filter(Product.name == product_name).first()
        if existing:
            skipped_count += 1
            continue
        
        try:
            # Map Excel columns to model fields
            product = Product(
                name=str(product_name),
                product_code=str(sheet.cell(row, 2).value) if sheet.cell(row, 2).value else None,
                min_batch_size=float(sheet.cell(row, 4).value or 0),
                max_batch_size=float(sheet.cell(row, 5).value or 0),
                ade_pde=float(sheet.cell(row, 6).value or 0),
                min_dose=float(sheet.cell(row, 7).value or 0),
                max_dose=float(sheet.cell(row, 8).value or 0),
                swab_recovery=float(sheet.cell(row, 9).value or 100),
                lod=float(sheet.cell(row, 10).value or 0.1),
                loq=float(sheet.cell(row, 11).value or 0.5),
                swab_dilution=float(sheet.cell(row, 12).value or 20),
                swab_surface_area=0.01,  # Default
                solubility=str(sheet.cell(row, 15).value or "Soluble"),
                hardest_to_clean=str(sheet.cell(row, 16).value or "Medium"),
                plant=str(sheet.cell(row, 17).value or "Plant-1")
            )
            db.add(product)
            imported_count += 1
        except Exception as e:
            skipped_count += 1
            errors.append(f"Row {row}: {str(e)}")
            continue
    
    db.commit()
    
    # Close the file
    file.file.close()
    
    return {
        "message": f"Import completed",
        "imported": imported_count,
        "skipped": skipped_count,
        "errors": errors[:10]  # Return first 10 errors
    }


def import_equipment_from_excel(file: UploadFile, db: Session) -> dict:
    """Import equipment from Excel file (SYNC VERSION)"""
    
    contents = file.file.read()
    workbook = openpyxl.load_workbook(io.BytesIO(contents))
    
    sheet = None
    for sheet_name in ["Equipment Details", "Equipment"]:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            break
    
    if not sheet:
        return {"error": "Could not find equipment sheet. Expected 'Equipment Details'"}
    
    imported_count = 0
    skipped_count = 0
    
    # Start from row 6 (data starts in original Excel)
    for row in range(6, sheet.max_row + 1):
        eq_name = sheet.cell(row, 2).value  # Column B
        if not eq_name:
            continue
        
        eq_id = sheet.cell(row, 3).value
        existing = db.query(Equipment).filter(Equipment.equipment_id == eq_id).first()
        if existing:
            skipped_count += 1
            continue
        
        try:
            equipment = Equipment(
                name=str(eq_name),
                equipment_id=str(eq_id),
                capacity=float(sheet.cell(row, 4).value) if sheet.cell(row, 4).value else None,
                surface_area=float(sheet.cell(row, 5).value or 0),
                used_for=str(sheet.cell(row, 6).value or ""),
                cleaning_procedure=str(sheet.cell(row, 8).value or ""),
                plant="Plant-1"  # Default
            )
            db.add(equipment)
            imported_count += 1
        except Exception as e:
            skipped_count += 1
            continue
    
    db.commit()
    
    # Close the file
    file.file.close()
    
    return {
        "message": f"Import completed",
        "imported": imported_count,
        "skipped": skipped_count
    }